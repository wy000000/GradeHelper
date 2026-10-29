"""报告解析：docx → Markdown + 图片，全部驻留内存。

- 主解析器：docling（版面分析，表格/图片/阅读顺序还原更准确）
- 兜底解析器：python-docx（docling 失败时使用）
- 文件名格式：课程名称-专业班级-学号-姓名.docx（课程名称可含 '-'，按右侧三段切分）
- md 中图片位置使用占位符 [[PIC:Pn]]，后续由 vllm 描述结果替换
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from .models import ParsedReport, ReportImage

logger = logging.getLogger(__name__)

DOCX_SUFFIX = ".docx"
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff"}
PIC_PATTERN = re.compile(r"\[\[PIC:P(\d+)]]")


def parse_filename(path: Path) -> tuple[str, str, str, str] | None:
    """按 '课程名称-专业班级-学号-姓名.docx' 解析文件名；不符合格式返回 None。"""
    if path.suffix.lower() != DOCX_SUFFIX or path.name.startswith("~$"):
        return None
    parts = path.stem.rsplit("-", 3)
    if len(parts) != 4:
        return None
    course, klass, student_id, name = parts
    if not all(p.strip() for p in (course, klass, student_id, name)):
        return None
    return course, klass, student_id, name


def scan_reports(reports_dir: str | Path) -> list[Path]:
    """扫描目录（含子目录），仅保留文件名格式正确的 docx。"""
    root = Path(reports_dir)
    if not root.is_dir():
        raise FileNotFoundError(f"报告目录不存在: {root}")
    return [p for p in sorted(root.rglob("*"))
            if p.is_file() and parse_filename(path=p) is not None]


def _sniff_ext(data: bytes) -> str | None:
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    if data[:3] == b"\xff\xd8\xff":
        return "jpg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    if data[:2] == b"BM":
        return "bmp"
    return None


def _content_key(data: bytes) -> str:
    import hashlib
    return hashlib.md5(data).hexdigest()


# ---------------------------------------------------------------- docling 主解析

def _extract_picture_bytes(doc, item) -> tuple[bytes | None, str | None]:
    """从 docling 文档工件中取图片二进制与文件名。"""
    art_no = getattr(item, "image_item_no", None)
    if art_no is None:
        return None, None
    try:
        artifact = doc.artifacts[art_no]
    except (IndexError, KeyError, TypeError):
        return None, None
    binary = getattr(artifact, "binary", None)
    if binary:
        return bytes(binary), _artifact_name(artifact)
    uri = getattr(artifact, "uri", None)
    if uri:
        try:
            return Path(uri).read_bytes(), Path(uri).name
        except OSError:
            pass
    return None, None


def _artifact_name(artifact) -> str:
    uri = getattr(artifact, "uri", None)
    return Path(uri).name if uri else ""


def _replace_image_ref(md: str, art_name: str, placeholder: str) -> str:
    """把 md 中引用 art_name 的图片语法替换为占位符。"""
    pattern = re.compile(
        r"!\[[^\]]*]\(\s*<?[^)\s]*" + re.escape(art_name) + r"[^)\s]*>?\s*\)",
        re.IGNORECASE,
    )
    new_md, n = pattern.subn(placeholder, md, count=1)
    if n == 0:  # 退化：仅替换文件名
        new_md = md.replace(art_name, placeholder, 1)
    return new_md


def _docling_parse(path: Path) -> tuple[str, list[ReportImage]]:
    """用 docling 解析 docx，返回 (md 文本, 图片列表)。md 中图片位置为占位符。

    docling 对 DOCX 默认不导出图片工件（Markdown 中为 `<!-- image -->` 注释），
    因此 Markdown 由 docling 生成，图片二进制在工件不可用时用 python-docx
    按正文顺序提取后依次回填占位位置。
    """
    from docling.document_converter import DocumentConverter

    doc = DocumentConverter().convert(str(path)).document
    md = doc.export_to_markdown() or ""

    # 1) 优先尝试 docling 自带图片工件（不同版本支持程度不同）
    pictures = sorted(
        doc.pictures,
        key=lambda it: ((getattr(it, "page_no", 0) or 0),
                        (getattr(getattr(it, "bbox", None), "t", 0) or 0)),
    )
    images: list[ReportImage] = []
    artifact_names: list[str] = []
    for item in pictures:
        data, art_name = _extract_picture_bytes(doc, item)
        if data is None:
            continue
        ext = _sniff_ext(data) or Path(art_name or "").suffix.lstrip(".") or "png"
        images.append(ReportImage(index=len(images) + 1, ext=ext, data=data))
        if art_name:
            artifact_names.append(art_name)

    # 2) 工件不可用时，用 python-docx 按正文出现顺序提取
    if not images:
        images = _extract_ordered_images(path)

    # 3) 把 md 中的图片位置替换为占位符 [[PIC:Pn]]
    for idx, art_name in enumerate(artifact_names, start=1):
        if idx <= len(images) and art_name and art_name in md:
            md = _replace_image_ref(md, art_name, f"[[PIC:P{images[idx - 1].index}]]")
    md = _place_placeholders(md, len(images))

    return md, images


def _extract_ordered_images(path: Path) -> list[ReportImage]:
    """python-docx 按正文出现顺序提取图片（内容去重，保持首现顺序）。"""
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(str(path))
    images: list[ReportImage] = []
    seen: set[str] = set()
    for blip in document.element.body.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in document.part.rels:
            continue
        try:
            blob = document.part.rels[rid].target_part.blob
        except (KeyError, OSError):
            continue
        key = _content_key(blob)
        if key in seen:
            continue
        seen.add(key)
        images.append(ReportImage(index=len(images) + 1,
                                  ext=_sniff_ext(blob) or "png", data=blob))
    if not images:  # 最终兜底：word/media 全量
        images = [ReportImage(index=i, ext=ext, data=data)
                  for i, (data, ext) in enumerate(_extract_media(path), start=1)]
    return images


def _place_placeholders(md: str, count: int) -> str:
    """按出现顺序把 md 中的图片位置替换为 [[PIC:Pn]]；未定位的占位符追加到文末。"""
    if count <= 0:
        return md
    n = 0

    def _sub_comment(_match: re.Match) -> str:
        nonlocal n
        n += 1
        return f"[[PIC:P{n}]]" if n <= count else ""

    # docling 的 `<!-- image -->` 注释占位
    md = re.sub(r"<!--\s*image\s*-->", _sub_comment, md, flags=re.IGNORECASE)
    if n < count:  # 剩余图片按 markdown 图片语法位置回填

        def _sub_mdimg(_match: re.Match) -> str:
            nonlocal n
            n += 1
            return f"[[PIC:P{n}]]"

        md = re.sub(r"!\[[^\]]*]\([^)]*\)", _sub_mdimg, md)

    total = len(PIC_PATTERN.findall(md))
    if total < count:  # 仍未定位的图片，占位符追加到文末
        md += "\n\n" + "\n".join(f"[[PIC:P{k}]]" for k in range(total + 1, count + 1))
    elif total > count:  # 多余占位符清理
        for k in range(count + 1, total + 1):
            md = md.replace(f"[[PIC:P{k}]]", "", 1)
    return md


# ------------------------------------------------------------ python-docx 兜底解析

def _looks_like_list(text: str) -> bool:
    stripped = text.lstrip()
    return bool(re.match(r"^([-*·•]|\d+[.、)]|[（(]\d+[)）])\s+\S", stripped))


def _table_to_md(table) -> list[str]:
    rows = []
    for r in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in r.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if len(rows) >= 2:
        rows.insert(1, "|" + "---|" * len(table.columns))
    return rows


def _collect_inline_images(para, images: list[ReportImage], seen: set[str]) -> None:
    """按行内 blip 关系提取图片 bytes，段落文本中追加 [[IMG:n]] 占位。"""
    from docx.oxml.ns import qn

    part = para.part
    for blip in para._p.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in part.rels:
            continue
        try:
            blob = part.rels[rid].target_part.blob
        except (KeyError, OSError):
            continue
        key = _content_key(blob)
        if key in seen:
            continue
        seen.add(key)
        ext = _sniff_ext(blob) or "png"
        images.append(ReportImage(index=len(images) + 1, ext=ext, data=blob))
        para.text += f"[[IMG:{len(images)}]]"  # 仅内存对象，不影响磁盘文件


def _extract_media(path: Path) -> list[tuple[bytes, str]]:
    """直接从 docx 压缩包的 word/media 提取全部图片（兜底补充）。"""
    import zipfile

    out: list[tuple[bytes, str]] = []
    try:
        with zipfile.ZipFile(path) as z:
            for name in sorted(z.namelist()):
                if name.startswith("word/media/"):
                    data = z.read(name)
                    sniffed = _sniff_ext(data)
                    ext = sniffed or Path(name).suffix.lstrip(".").lower() or "png"
                    if f".{ext}" in IMAGE_EXTS or sniffed:
                        out.append((data, ext))
    except (OSError, zipfile.BadZipFile):
        pass
    return out


def _docx_fallback_parse(path: Path) -> tuple[str, list[ReportImage]]:
    """python-docx 兜底：段落/表格转 md，图片按行内关系 + word/media 提取。"""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []
    images: list[ReportImage] = []
    seen: set[str] = set()

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, document)
            _collect_inline_images(para, images, seen)
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            style = (para.style.name or "") if para.style is not None else ""
            if style.lower().startswith("heading"):
                try:
                    level = min(int(style.split()[-1]), 6)
                except (ValueError, IndexError):
                    level = 2
                lines.append("#" * level + " " + text)
            elif _looks_like_list(text):
                lines.append("- " + text.lstrip("-*·• \t"))
            else:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            lines.extend(_table_to_md(Table(child, document)))

    # word/media 中未通过行内关系取到的图片，按压缩包内顺序补充
    for data, ext in _extract_media(path):
        if _content_key(data) in seen:
            continue
        images.append(ReportImage(index=len(images) + 1, ext=ext, data=data))

    md = "\n\n".join(lines)
    for img in images:
        md = md.replace(f"[[IMG:{img.index}]]", f"[[PIC:P{img.index}]]", 1)
    located = set(int(m) for m in PIC_PATTERN.findall(md))
    missing = [img.index for img in images if img.index not in located]
    if missing:
        md += "\n\n" + "\n".join(f"[[PIC:P{n}]]" for n in missing)
    return md, images


# ------------------------------------------------------------------------ 对外接口

def parse_report(path: Path) -> ParsedReport:
    """解析单份报告；docling 优先，失败回退 python-docx。"""
    meta = parse_filename(path)
    if meta is None:
        raise ValueError(f"文件名不符合 '课程名称-专业班级-学号-姓名.docx' 格式: {path.name}")
    course, klass, student_id, name = meta

    warnings: list[str] = []
    try:
        md, images = _docling_parse(path)
    except Exception as exc:  # docling 依赖模型，任何异常都走兜底
        warnings.append(f"docling 解析失败，回退 python-docx: {exc}")
        logger.warning("%s: docling 解析失败(%s)，回退 python-docx", path.name, exc)
        md, images = _docx_fallback_parse(path)

    return ParsedReport(course=course, klass=klass, student_id=student_id, name=name,
                        source_path=str(path), md_text=md, images=images, warnings=warnings)


def parse_all(reports_dir: str | Path) -> list[ParsedReport]:
    """扫描并解析目录下所有合规报告，解析结果全部驻留内存。"""
    reports = []
    for path in scan_reports(reports_dir):
        try:
            reports.append(parse_report(path))
        except Exception as exc:
            logger.error("解析失败 %s: %s", path.name, exc)
    return reports


def embed_captions(md_text: str, images: list[ReportImage], template: str) -> str:
    """把图片描述嵌入 md 中对应占位符位置。"""
    result = md_text
    for img in images:
        if not img.caption:
            continue
        block = template.format(n=img.index, caption=img.caption.strip())
        result = result.replace(f"[[PIC:P{img.index}]]", block, 1)
    result = PIC_PATTERN.sub("", result)  # 清理未替换的占位符
    return re.sub(r"\n{3,}", "\n\n", result).strip() + "\n"
