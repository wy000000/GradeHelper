"""总编排：doc(解析+程序规则) → pic(vllm识图) → sen(llm语义评分)，可独立执行。"""

from __future__ import annotations

import logging
import re
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from .config import AppConfig
from .excel_io import (
    read_result_scores, read_rules, write_results, build_expected_headers, RESULT_SHEET
)
from openpyxl import load_workbook
from .image_caption import caption_images, save_captions_txt
from .llm_client import LLMClient
from .models import ItemScore, ParsedReport, ReportImage, RuleItem
from .parser import embed_captions, parse_all
from .rules_engine import run_program_rules

logger = logging.getLogger(__name__)


@dataclass
class StageLogEntry:
    """单个阶段的运行日志记录。"""

    stage: str                      # 阶段名：doc / pic / sen
    start_time: str                 # 开始时间 "YYYY-MM-DD HH:MM:SS"
    end_time: str                   # 结束时间
    duration_seconds: float         # 用时（秒）
    processed_count: int            # 共处理多少份
    start_student_id: str | None    # 开始学号（第一份报告的学号）
    end_student_id: str | None      # 结束学号（最后一份报告的学号）
    error: str | None = None        # 出错原因（若无则为 None）


class Grader:
    def __init__(self, config: AppConfig):
        self.config = config
        self.stage_logs: list[StageLogEntry] = []
        self._log_path: Path | None = None

    def _resolve_log_path(self, template: str | Path) -> Path:
        """确定运行日志文件路径。

        - 若 config.paths.log_path 非空，直接使用；
        - 否则与模板文件同目录，文件名为 <模板名>.log。
        """
        cfg = self.config
        log_path = (cfg.paths.log_path or "").strip()
        if log_path:
            return Path(log_path)
        return Path(str(template) + ".log")

    def _write_run_log(self) -> None:
        """把所有阶段的运行日志写入日志文件。"""
        if not self._log_path:
            return
        log_path = Path(self._log_path)
        log_path.parent.mkdir(parents=True, exist_ok=True)

        lines: list[str] = []
        lines.append("=" * 60)
        lines.append("GradeHelper 运行日志")
        lines.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("=" * 60)
        lines.append("")

        for entry in self.stage_logs:
            lines.append(f"阶段: {entry.stage}")
            lines.append(f"开始时间: {entry.start_time}")
            lines.append(f"结束时间: {entry.end_time}")
            lines.append(f"用时: {entry.duration_seconds:.2f} 秒")
            lines.append(f"共处理: {entry.processed_count} 份")
            if entry.start_student_id:
                lines.append(f"开始学号: {entry.start_student_id}")
            if entry.end_student_id:
                lines.append(f"结束学号: {entry.end_student_id}")
            if entry.error:
                lines.append(f"状态: 失败 - {entry.error}")
            else:
                lines.append("状态: 成功")
            lines.append("-" * 60)
            lines.append("")

        log_path.write_text("\n".join(lines), encoding="utf-8")
        logger.info("运行日志已写入: %s", log_path)

    # ------------------------------------------------------------------ 阶段 1：doc（解析 + 程序规则打分）

    def stage_doc(self, reports_dir: str | Path, template: str | Path,
                  skip_images: int = 2) -> None:
        """阶段 doc：解析所有文档 + 程序化规则打分，写入 result 表（LLM 项先写 0）。"""
        cfg = self.config
        reports_dir = Path(reports_dir)
        t0 = time.perf_counter()
        start_dt = datetime.now()
        reports: list[ParsedReport] = []
        error: str | None = None
        try:
            logger.info("【阶段 doc】解析报告目录: %s", reports_dir)
            reports = parse_all(reports_dir)
            if not reports:
                raise RuntimeError(f"目录中没有符合命名格式的报告: {reports_dir}")

            # 移除被 skip 的图片（不保存图片文件）及其占位符，保持剩余图片原始编号
            for report in reports:
                before = len(report.images)
                report.images = [img for img in report.images if img.index > skip_images]
                for skip_idx in range(1, skip_images + 1):
                    report.md_text = report.md_text.replace(f"[[PIC:P{skip_idx}]]", "")
                logger.info("  已解析 %s（图片 %d 张，跳过前 %d 张，剩余 %d 张）",
                            report.base_name, before, skip_images, len(report.images))

            # 始终保存解析产物（md 与图片）到与原文件相同路径
            self._save_parsed(reports)

            # 读取评分要求（rule 表）
            logger.info("【阶段 doc】读取评分模板 rule 表: %s", template)
            rule_items = read_rules(template)

            # 程序化 rule 项逐项批处理
            logger.info("【阶段 doc】程序化规则评分（逐 rule 批处理）")
            scores, warnings = run_program_rules(reports, rule_items)
            for w in warnings:
                logger.warning("  %s", w)

            # LLM 项分数先写 0，sen 阶段再更新
            llm_rule_keys = [r.key for r in rule_items if r.is_llm]
            for sid_items in scores.values():
                for key in llm_rule_keys:
                    sid_items[key] = ItemScore(key=key, score=0.0, confidence=None, comment="")

            # 写入 result 表（按学号匹配：存在则更新，不存在则添加）
            write_results(template, template, reports, rule_items, scores, mode="upsert")
            logger.info("【阶段 doc】完成，程序化规则分数已写入 result 表（LLM 项已预置 0），耗时 %.2f 秒",
                        time.perf_counter() - t0)
        except Exception as exc:  # noqa: BLE001
            error = str(exc)
            raise
        finally:
            end_dt = datetime.now()
            elapsed = time.perf_counter() - t0
            start_sid = reports[0].student_id if reports else None
            end_sid = reports[-1].student_id if reports else None
            self.stage_logs.append(StageLogEntry(
                stage="doc",
                start_time=start_dt.strftime("%Y-%m-%d %H:%M:%S"),
                end_time=end_dt.strftime("%Y-%m-%d %H:%M:%S"),
                duration_seconds=round(elapsed, 2),
                processed_count=len(reports),
                start_student_id=start_sid,
                end_student_id=end_sid,
                error=error,
            ))

    # ------------------------------------------------------------------ 阶段 2：pic（vllm 图片识别 + 更新 md）

    def stage_pic(self, reports_dir: str | Path, template: str | Path,
                  skip_images: int = 2) -> None:
        """阶段 pic：vllm 识别所有图片，将描述嵌入 md 文档并落盘。"""
        cfg = self.config
        reports_dir = Path(reports_dir)
        t0 = time.perf_counter()
        start_dt = datetime.now()
        reports: list[ParsedReport] = []
        error: str | None = None
        try:
            logger.info("【阶段 pic】从已保存的 md 与图片加载报告: %s", reports_dir)
            reports = self._load_reports_for_pic(reports_dir, skip_images)
            if not reports:
                raise RuntimeError(f"目录中没有找到已保存的 md 文件: {reports_dir}")

            llm_rules_flag = any(r.is_llm for r in read_rules(template))
            if llm_rules_flag:
                logger.info("【阶段 pic】vllm 批量识别图片")
                caption_images(reports, cfg.vllm, cfg.prompts.image_prompt)
                saved = save_captions_txt(reports, cfg.paths.captions_dir or None)
                logger.info("  图片描述已另存 %d 个 txt", saved)
            for report in reports:
                report.md_text = embed_captions(report.md_text, report.images,
                                                cfg.prompts.image_context)

            # 保存更新后的 md（带图片描述）
            self._save_parsed(reports)
            logger.info("【阶段 pic】完成，图片描述已嵌入 md 并落盘，耗时 %.2f 秒",
                        time.perf_counter() - t0)
        except Exception as exc:  # noqa: BLE001
            error = str(exc)
            raise
        finally:
            end_dt = datetime.now()
            elapsed = time.perf_counter() - t0
            start_sid = reports[0].student_id if reports else None
            end_sid = reports[-1].student_id if reports else None
            self.stage_logs.append(StageLogEntry(
                stage="pic",
                start_time=start_dt.strftime("%Y-%m-%d %H:%M:%S"),
                end_time=end_dt.strftime("%Y-%m-%d %H:%M:%S"),
                duration_seconds=round(elapsed, 2),
                processed_count=len(reports),
                start_student_id=start_sid,
                end_student_id=end_sid,
                error=error,
            ))

    # ------------------------------------------------------------------ 阶段 3：sen（llm 语义评分 + 更新 result 表）

    def stage_sen(self, reports_dir: str | Path, template: str | Path,
                  output: str | Path, skip_images: int = 2) -> Path:
        """阶段 sen：llm 语义评分（单文件批量），写入最终 result 表。"""
        cfg = self.config
        reports_dir = Path(reports_dir)
        t0 = time.perf_counter()
        start_dt = datetime.now()
        reports: list[ParsedReport] = []
        error: str | None = None
        try:
            # 读取 result 表中已存在的所有评分（不限 LLM 项，避免覆盖程序化规则分数）
            rule_items = read_rules(template)
            llm_rules = [r for r in rule_items if r.is_llm]
            existing_scores = read_result_scores(template)  # 读取全部评分列

            # 从磁盘读取已更新的 md（pic 阶段保存的）
            reports = self._load_reports_for_pic(reports_dir, skip_images)

            if not llm_rules:
                logger.info("【阶段 sen】无 LLM 评分项，沿用已有分数")
                final_scores: dict[str, dict[str, ItemScore]] = {}
                for report in reports:
                    sid = report.student_id
                    row: dict[str, ItemScore] = {}
                    for rule in rule_items:
                        prev = existing_scores.get(sid, {}).get(rule.key, 0.0)
                        row[rule.key] = ItemScore(key=rule.key, score=prev,
                                                  confidence=None, comment="")
                    final_scores[sid] = row
            else:
                # sen 阶段前置检查：result 表表头必须与当前规则完全匹配，
                # 不匹配则报错退出，不进行任何 LLM 提交。
                expected_header = build_expected_headers(rule_items)
                wb_check = load_workbook(template, data_only=False)
                if RESULT_SHEET in wb_check.sheetnames:
                    ws_check = wb_check[RESULT_SHEET]
                    existing_header = (
                        [cell.value for cell in ws_check[1]] if ws_check.max_row >= 1 else []
                    )
                else:
                    existing_header = []
                if existing_header != expected_header:
                    raise RuntimeError(
                        "sen 阶段：result 表表头与当前评分规则不匹配，程序退出，未进行 LLM 提交。\n"
                        f"期望表头：{expected_header}\n"
                        f"现有表头：{existing_header}"
                    )

                logger.info("【阶段 sen】LLM 评分项批量评估（每份报告一次提交 %d 项）", len(llm_rules))
                client = LLMClient(cfg.llm)
                final_scores = {}
                total = len(reports)
                for idx, report in enumerate(reports, 1):
                    sid = report.student_id
                    row: dict[str, ItemScore] = {}
                    # 先把程序化规则分数填进来
                    for rule in rule_items:
                        if not rule.is_llm:
                            prev = existing_scores.get(sid, {}).get(rule.key, 0.0)
                            row[rule.key] = ItemScore(key=rule.key, score=prev, comment="")
                    # 批量 LLM 评分（一次提交所有 LLM 项）；失败抛出异常，由 main 捕获后退出程序
                    batch_scores = self._grade_with_llm_batch(client, report, llm_rules)
                    for rule in llm_rules:
                        row[rule.key] = batch_scores[rule.key]
                    final_scores[sid] = row
                    logger.info("【阶段 sen】%d/%d, ID %s：", idx, total, sid)

            # 汇总写入 result 表（按学号匹配：存在则更新，不存在则报错退出）
            out = write_results(template, output, reports, rule_items, final_scores, mode="strict")
            logger.info("【阶段 sen】完成，最终评分已写入: %s，耗时 %.2f 秒", out,
                        time.perf_counter() - t0)
            return out
        except Exception as exc:  # noqa: BLE001
            error = str(exc)
            raise
        finally:
            end_dt = datetime.now()
            elapsed = time.perf_counter() - t0
            start_sid = reports[0].student_id if reports else None
            end_sid = reports[-1].student_id if reports else None
            self.stage_logs.append(StageLogEntry(
                stage="sen",
                start_time=start_dt.strftime("%Y-%m-%d %H:%M:%S"),
                end_time=end_dt.strftime("%Y-%m-%d %H:%M:%S"),
                duration_seconds=round(elapsed, 2),
                processed_count=len(reports),
                start_student_id=start_sid,
                end_student_id=end_sid,
                error=error,
            ))

    # ------------------------------------------------------------------ 总流程调度

    def run(self, reports_dir: str | None = None, template: str | None = None,
            output: str | None = None, stage: str = "full",
            skip_images: int = 2) -> Path | None:
        cfg = self.config
        reports_dir = reports_dir or cfg.paths.reports_dir
        template = template or cfg.paths.template
        output = output or cfg.paths.output or template

        # 确定运行日志文件路径
        self._log_path = self._resolve_log_path(template)

        try:
            if stage == "full":
                self.stage_doc(reports_dir, template, skip_images)
                self.stage_pic(reports_dir, template, skip_images)
                return self.stage_sen(reports_dir, template, output, skip_images)

            if stage == "doc":
                self.stage_doc(reports_dir, template, skip_images)
                return None

            if stage == "pic":
                self.stage_pic(reports_dir, template, skip_images)
                return None

            if stage == "sen":
                return self.stage_sen(reports_dir, template, output, skip_images)

            raise ValueError(f"未知的 stage 参数: {stage}")
        finally:
            # 无论成功或出错，都写入运行日志
            self._write_run_log()

    # ------------------------------------------------------------------ 从磁盘加载报告（供 pic 阶段使用，不重新解析 word）

    PIC_PATTERN = re.compile(r"\[\[PIC:P(\d+)\]\]")

    def _load_reports_for_pic(self, reports_dir: Path, skip_images: int) -> list[ParsedReport]:
        """从 doc 阶段已保存的 md 与图片文件加载报告，不重新解析 word。

        扫描目录下所有 *.md 文件，从文件名解析元数据，读取 md 内容，
        再从 md 中提取 [[PIC:Pn]] 占位符，按 {base_name}-P{n}.{ext} 模式
        查找并读取对应图片文件。
        """
        IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff"}
        reports: list[ParsedReport] = []
        for md_path in sorted(reports_dir.glob("*.md")):
            if md_path.name.startswith("~$"):
                continue
            try:
                # 从 md 文件名解析：课程名称-专业班级-学号-姓名.md
                parts = md_path.stem.rsplit("-", 3)
                if len(parts) != 4:
                    logger.warning("跳过不符合命名格式的 md 文件: %s", md_path.name)
                    continue
                course, klass, student_id, name = parts
                if not all(p.strip() for p in (course, klass, student_id, name)):
                    continue

                md_text = md_path.read_text(encoding="utf-8")
                base_name = md_path.stem

                # 从 md 中提取所有 [[PIC:Pn]] 占位符，确定图片索引
                pic_indices = sorted({int(m) for m in self.PIC_PATTERN.findall(md_text)})

                images: list[ReportImage] = []
                for pic_idx in pic_indices:
                    img_data = None
                    img_ext = None
                    for ext in IMAGE_EXTS:
                        candidate = md_path.parent / f"{base_name}-P{pic_idx}{ext}"
                        if candidate.is_file():
                            img_data = candidate.read_bytes()
                            img_ext = ext.lstrip(".")
                            break
                    if img_data is None:
                        logger.warning("未找到图片文件: %s-P%d.*", base_name, pic_idx)
                        continue
                    images.append(ReportImage(index=pic_idx, ext=img_ext, data=img_data))

                # 应用跳图策略：只保留 index > skip_images 的图片进行处理，
                # 但不改变图片原始编号（始终反映在文档中的实际位置）
                images = [img for img in images if img.index > skip_images]

                reports.append(ParsedReport(
                    course=course, klass=klass, student_id=student_id, name=name,
                    source_path=str(md_path), md_text=md_text, images=images,
                ))
            except Exception as exc:  # noqa: BLE001
                logger.error("加载报告失败 %s: %s", md_path.name, exc)
        return reports

    # ------------------------------------------------------------------ LLM 批量评估（单文件一次性提交所有 LLM 项）

    def _grade_with_llm_batch(self, client: LLMClient, report: ParsedReport,
                              llm_rules: list[RuleItem]) -> dict[str, ItemScore]:
        """对单个文件的所有 LLM 评分项一次性提交，返回 {key: ItemScore}。

        失败不自行捕获，异常向上抛出。client.chat_json 内部已按 cfg.max_retries
        重试，耗尽后会抛出 RuntimeError，由 stage_sen / main 决定退出程序。
        """
        items_section: list[str] = []
        for rule in llm_rules:
            lo, hi = rule.score_range()
            score_range = f"{lo:g}-{hi:g}" if lo is not None and hi is not None else "按评分标准"
            items_section.append(f"- {rule.key}：{rule.evaluation}，分数范围 {score_range}。")
        items_text = "\n".join(items_section)

        prompt = self.config.prompts.llm_batch.format(
            items_section=items_text, report=report.md_text,
        )
        data = client.chat_json([{"role": "user", "content": prompt}])

        result: dict[str, ItemScore] = {}
        for rule in llm_rules:
            entry = data.get(rule.key) or {}
            score = _to_float(entry.get("score"))
            if score is None:
                lo, _ = rule.score_range()
                score = lo if lo is not None else 0.0
            confidence = _to_float(entry.get("confidence"))
            comment = str(entry.get("comment") or rule.comment or "").strip()
            result[rule.key] = ItemScore(key=rule.key, score=score,
                                         confidence=confidence, comment=comment)
            logger.info("  %s / %s -> score=%s confidence=%s (batch)",
                        report.student_id, rule.key, score, confidence)
        return result

    # ------------------------------------------------------------------ 可选落盘

    @staticmethod
    def _save_parsed(reports: list[ParsedReport]) -> None:
        """把解析出的 md 与图片落盘到与原文件相同路径。"""
        for report in reports:
            target_dir = Path(report.source_path).parent
            (target_dir / f"{report.base_name}.md").write_text(report.md_text, encoding="utf-8")
            for image in report.images:
                (target_dir / f"{report.base_name}-P{image.index}.{image.ext}").write_bytes(image.data)
        logger.info("解析产物（md/图片）已落盘 %d 份报告", len(reports))


def _to_float(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None